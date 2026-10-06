"""Gera um .pptx REAL (python-pptx) para os testes de validação de upload.
Uso: python3 -I generate.py <saida.pptx>   (sai com 0 e escreve o arquivo; sai com 2 se python-pptx não estiver instalado)."""
import sys

try:
    from pptx import Presentation
    from pptx.util import Inches
    import struct, zlib
except ImportError:
    sys.exit(2)

def tiny_png(path):
    # PNG 4x4 laranja, sem depender de Pillow para criar (o python-pptx lê com Pillow, que ele já exige)
    w = h = 4
    raw = b''.join(b'\x00' + b'\xf7\x8c\x16' * w for _ in range(h))
    def chunk(t, d):
        c = struct.pack('>I', len(d)) + t + d
        return c + struct.pack('>I', zlib.crc32(t + d) & 0xffffffff)
    data = b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b'')
    open(path, 'wb').write(data)

out = sys.argv[1]
prs = Presentation()
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = 'Apresentação de teste'
s.placeholders[1].text = 'Gerada por python-pptx'
s2 = prs.slides.add_slide(prs.slide_layouts[5])
s2.shapes.title.text = 'Slide com imagem'
png = out + '.png'
tiny_png(png)
s2.shapes.add_picture(png, Inches(1), Inches(2), Inches(2), Inches(2))
prs.save(out)
